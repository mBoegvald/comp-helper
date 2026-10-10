#!/usr/bin/env python3
"""One-off move from the Excel workbooks (and their CSVs) into data/pickhelper.db. Needs openpyxl.

  python3 migrate_xlsx.py            # refuses if the database already has hand data
  python3 migrate_xlsx.py --force    # replace the hand layer with what the workbooks hold

Imports, per role: data/lolalytics_<role>.csv -> lola, workbook champions + data/champions_<role>.txt -> pool,
data/reddit_tips.csv -> reddit_tips/reddit_snippet, and from the workbook only what differs from what an update
would generate: champion fields that differ from role_data.py, the hand part of Good into / Struggles into
(the text after the data-derived names), Results that differ from the Lolalytics label, lane tips, and the
Comps/Notes sheets when they are not the generated defaults.
"""
import csv
import json
import sys
from pathlib import Path

import openpyxl

import db
import role_data

HERE = Path(__file__).resolve().parent
OLD_XLSX = {"top": "roles/top.xlsx", "jungle": "roles/jungle.xlsx", "mid": "midlane_overview.xlsx",
            "bot": "roles/bot.xlsx", "support": "roles/support.xlsx"}
DEFAULT_COMPS = [["Style", "Words used in 'Strong in comps'"], ["wombo", "wombo, teamfight"], ["poke", "poke, siege"],
                 ["pick", "pick"], ["dive", "dive"], ["split", "split, flank"], ["scaling", "scaling, late, front-to-back"]]


def num(x, cast=float):
    return cast(float(x)) if x not in (None, "") else None


def import_lola(conn, role):
    path = HERE / "data" / f"lolalytics_{role}.csv"
    rows = list(csv.DictReader(path.open(encoding="utf-8")))
    conn.execute("DELETE FROM lola WHERE role = ?", (role,))
    conn.executemany("INSERT OR REPLACE INTO lola VALUES (?,?,?,?,?,?,?,?,?,?,?,?)", [
        (role, r["champion"], r["opponent"], num(r["wr"]), num(r["delta_vs_avg"]), num(r["delta_norm"]),
         num(r["opp_avg_wr"]), num(r["games"], int), r["patch"], r["tier"], r["lane"], r["fetched_at"]) for r in rows])
    return len(rows)


def import_tips(conn):
    path = HERE / "data" / "reddit_tips.csv"
    conn.execute("DELETE FROM reddit_tips"); conn.execute("DELETE FROM reddit_snippet")
    n = 0
    for r in csv.DictReader(path.open(newline="", encoding="utf-8")):
        conn.execute("INSERT INTO reddit_tips VALUES (?,?,?,?,?)",
                     (r["champion"], r["opponent"], num(r["mentions"], int), num(r["weighted"]), r["newest"]))
        for i in (1, 2, 3):
            if r.get(f"tip{i}"):
                date, _, url = (r.get(f"src{i}") or "").partition(" ")
                conn.execute("INSERT INTO reddit_snippet VALUES (?,?,?,?,?,?)", (r["champion"], r["opponent"], i, r[f"tip{i}"], date, url))
        n += 1
    return n


def hand_note(text, data_list, default):
    """The hand part of a Good into / Struggles into cell ('<data names>; <hand text>'), or None if it is the default."""
    text, data = str(text or "").strip(), ", ".join(data_list)
    if text == data:
        note, ok = "", True
    elif data and text.startswith(data + "; "):
        note, ok = text[len(data) + 2:], True
    elif not data and text.startswith("; "):
        note, ok = text[2:], True
    else:
        note, ok = text, not data  # data names did not match: keep the whole cell as hand text
    note = note.strip()
    return (None if note in ("", default) else note), ok


def import_workbook(conn, role, report):
    wb = openpyxl.load_workbook(HERE / OLD_XLSX[role], read_only=True)
    comb = db.combined(db.lola_rows(conn, role))
    archs = role_data.ARCHETYPES.get(role, {})
    names = set()
    txt = HERE / "data" / f"champions_{role}.txt"
    if txt.exists():
        names |= {c.strip() for c in txt.read_text(encoding="utf-8").splitlines() if c.strip()}

    rows = list(wb["Champions"].iter_rows(values_only=True))
    hdr = list(rows[0])
    for r in rows[1:]:
        if not r or not r[0]:
            continue
        x = dict(zip(hdr, r))
        c = str(r[0]).strip()
        names.add(c)
        arch_d, dmg_d, comps_d, when_d, blind_d = db.defaults(role, c)
        arch = x.get("Archetype") or arch_d
        if arch == "Flex":  # placeholder: the role_data entry wins, as on every rebuild
            arch = arch_d
        good, bad = db.data_lists(comb, c)
        _w, g_def, b_def = archs.get(arch, ([], "", ""))
        g_note, ok1 = hand_note(x.get("Good into"), good, g_def)
        b_note, ok2 = hand_note(x.get("Struggles into"), bad, b_def)
        if not (ok1 and ok2):
            report["unparsed"].append(f"{role}/{c}")
        diff = lambda v, d: (str(v).strip() if v not in (None, "") and str(v).strip() != str(d) else None)
        vals = {"archetype": diff(arch, arch_d), "damage": diff(x.get("Damage"), dmg_d), "comps": diff(x.get("Strong in comps"), comps_d),
                "good_into": g_note, "struggles_into": b_note, "pick_when": diff(x.get("Pick when"), when_d),
                "blind_safe": diff(x.get("Blind-safe"), blind_d)}
        if any(vals.values()):
            conn.execute(f"INSERT OR REPLACE INTO hand_champion (role, champion, {', '.join(vals)}, updated_at) VALUES (?, ?, {', '.join('?' * len(vals))}, ?)",
                         (role, c, *vals.values(), db.now()))
            report["champions"] += 1
    conn.executemany("INSERT OR IGNORE INTO pool VALUES (?, ?)", [(role, c) for c in sorted(names)])

    rows = list(wb["Matchups"].iter_rows(values_only=True))
    h = [str(v).strip().lower() if v else "" for v in rows[0]]
    ti, ri, li = h.index("lane tip"), h.index("result for champion"), h.index("lola label")
    for r in rows[1:]:
        if not r or not r[0] or not r[1]:
            continue
        res, tip, lol = r[ri], r[ti], r[li]
        hand_res = res if res and str(res) != str(lol or "") else None
        if hand_res or tip:
            conn.execute("INSERT OR REPLACE INTO hand_matchup VALUES (?,?,?,?,?,?)",
                         (role, str(r[0]).strip(), str(r[1]).strip(), hand_res, tip or None, db.now()))
            report["labels"] += bool(hand_res); report["tips"] += bool(tip)

    for sheet in ("Comps", "Notes"):
        if sheet not in wb.sheetnames:
            continue
        cells = [[v for v in r] for r in wb[sheet].iter_rows(values_only=True) if r and any(v is not None for v in r)]
        cells = [c for c in cells if not str(c[0] or "").startswith("Built ")]
        if sheet == "Comps" and [[str(v) for v in c if v is not None] for c in cells] == DEFAULT_COMPS:
            continue
        for i, c in enumerate(cells):
            while c and c[-1] is None:
                c.pop()
            conn.execute("INSERT OR REPLACE INTO role_note VALUES (?,?,?,?)", (role, sheet, i, json.dumps(c, ensure_ascii=False)))
            report["notes"] += 1


def main():
    conn = db.connect()
    has_hand = conn.execute("SELECT (SELECT count(*) FROM hand_champion) + (SELECT count(*) FROM hand_matchup)").fetchone()[0]
    if has_hand and "--force" not in sys.argv:
        sys.exit(f"{db.PATH.name} already has {has_hand} hand rows; rerun with --force to replace them from the workbooks.")
    report = {"champions": 0, "labels": 0, "tips": 0, "notes": 0, "unparsed": []}
    with conn:
        for t in ("hand_champion", "hand_matchup", "role_note"):
            conn.execute(f"DELETE FROM {t}")
        for role in role_data.ROLES:
            n = import_lola(conn, role)
            print(f"{role}: {n} Lolalytics rows")
            db.touch(conn, f"lola_updated:{role}")
        print(f"reddit tips: {import_tips(conn)} pairs")
        db.touch(conn, "tips_updated")
        for role in role_data.ROLES:
            import_workbook(conn, role, report)
    print(f"hand layer: {report['champions']} champion rows, {report['labels']} labels, {report['tips']} lane tips, "
          f"{report['notes']} note rows")
    if report["unparsed"]:
        print(f"kept whole Good/Struggles text (data names did not match) for: {', '.join(report['unparsed'])}")
    conn.execute("VACUUM")


if __name__ == "__main__":
    main()
