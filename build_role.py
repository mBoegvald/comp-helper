#!/usr/bin/env python3
"""Build or refresh the workbook for a role (top, jungle, bot, support) from Lolalytics data + role_data.py.

  python3 build_role.py top            # -> roles/top.xlsx, data/lolalytics_top.csv, data/champions_top.txt
  python3 build_role.py all            # every role except mid (mid keeps midlane_overview.xlsx)
  python3 build_role.py top --from-csv # rebuild the workbook from data/lolalytics_top.csv without refetching

Champion pool: every champion with >= POOL_MIN_GAMES games against one of the role's seed champions on Lolalytics
(Emerald+, current patch), plus every champion listed for the role in role_data.CHAMPS.
Champions sheet: archetype, damage, comps, pick-when and blind-safe from role_data (fallback: 'Flex' archetype);
'Good into' / 'Struggles into' = the best / worst matchups by normalised delta plus the archetype's default text.
Matchups sheet: one row per champion/opponent pair with >= MIN_GAMES games, labelled from the normalised delta
exactly like fetch_lolalytics.py apply (both directions averaged), Lane tip empty for extract_tips.py to fill.
Rebuilding an existing workbook keeps any hand-edited Champions text and Lane tips and only refreshes the numbers.
"""
import csv
import datetime as dt
import sys
import time
from pathlib import Path

import openpyxl

import fetch_lolalytics as lola
import role_data

POOL_MIN_GAMES = 150
MIN_GAMES = 100  # rows kept; fetch_lolalytics.label marks < 200 as 'low sample'
TOP_N = 5
HERE = Path(__file__).resolve().parent


def pool_for(role: str, delay: float):
    lane = role_data.ROLES[role]["lane"]
    pool = {}
    for seed in role_data.ROLES[role]["seeds"]:
        for r in lola.fetch(seed, delay, lane):
            if r["games"] and r["games"] >= POOL_MIN_GAMES:
                pool[r["opponent"]] = max(pool.get(r["opponent"], 0), r["games"])
    for c in role_data.CHAMPS.get(role, {}):
        pool.setdefault(c, 0)
    return sorted(pool)


def fetch_all(role: str, champs, delay: float):
    lane = role_data.ROLES[role]["lane"]
    out = HERE / "data" / f"lolalytics_{role}.csv"
    out.parent.mkdir(exist_ok=True)
    now = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%MZ")
    rows = []
    for c in champs:
        for r in lola.fetch(c, delay, lane):
            rows.append({"champion": c, "fetched_at": now, "lane": lane, **r})
        print(f"  {c}: {sum(1 for r in rows if r['champion'] == c)} opponents", flush=True)
    with out.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=lola.CSV_FIELDS)
        w.writeheader()
        w.writerows(rows)
    print(f"wrote {len(rows)} rows to {out}")
    return rows


def combined(rows):
    """(champ, opp) -> dict(wr, dnorm, games, patch, tier, source) using both directions, like fetch_lolalytics apply."""
    direct = {(r["champion"], r["opponent"]): r for r in rows}
    out = {}
    for (c, o), fwd in direct.items():
        rev = direct.get((o, c))
        wrs, dns, games = [float(fwd["wr"])], [float(fwd["delta_norm"])], [fwd["games"] or 0]
        if rev:
            wrs.append(100 - float(rev["wr"])); dns.append(-float(rev["delta_norm"])); games.append(rev["games"] or 0)
        out[(c, o)] = {"wr": sum(wrs) / len(wrs), "dnorm": sum(dns) / len(dns), "games": min(games),
                       "patch": fwd["patch"], "tier": fwd["tier"], "source": "both" if rev else "direct"}
    return out


def norm_label(x) -> str:
    x = str(x or "").strip().lower()
    if x.startswith("fav"): return "Favored"
    if x.startswith("unfav"): return "Unfavored"
    if x.startswith("even") or x.startswith("skill") or x.startswith("same"): return "Even"
    return str(x or "")


def read_old(xlsx: Path):
    """Everything hand-made in an existing workbook: champion text, hand labels, lane tips, Comps and Notes."""
    old = {"champs": {}, "labels": {}, "tips": {}, "sheets": {}}
    if not xlsx.exists():
        return old
    wb0 = openpyxl.load_workbook(xlsx, read_only=True)
    rs = list(wb0["Champions"].iter_rows(values_only=True))
    old["champs"] = {str(r[0]).strip(): dict(zip(rs[0], r)) for r in rs[1:] if r and r[0]}
    rs = list(wb0["Matchups"].iter_rows(values_only=True))
    h = [str(x).strip().lower() if x else "" for x in rs[0]]
    def col(prefix):
        return next((i for i, x in enumerate(h) if x.startswith(prefix)), None)
    ti, ri, li = col("lane tip"), col("result"), col("lola label")
    for r in rs[1:]:
        if not r or not r[0] or not r[1]:
            continue
        key = (str(r[0]).strip(), str(r[1]).strip())
        if ti is not None and r[ti]:
            old["tips"][key] = r[ti]
        res = r[ri] if ri is not None else None
        lol = r[li] if li is not None else None
        # a result that is not just the copied Lolalytics label is a hand label worth keeping
        if res and (not lol or lol in ("no data",) or norm_label(res) != norm_label(lol) or str(res) != str(lol)):
            old["labels"][key] = res
    for name in ("Comps", "Notes"):
        if name in wb0.sheetnames:
            old["sheets"][name] = [list(r) for r in wb0[name].iter_rows(values_only=True) if r and any(x is not None for x in r)]
    return old


def build(role: str, delay: float, from_csv: bool = False):
    csv_path = HERE / "data" / f"lolalytics_{role}.csv"
    xlsx = HERE / role_data.ROLES[role]["xlsx"]
    xlsx.parent.mkdir(exist_ok=True)
    old = read_old(xlsx)
    champs_txt = HERE / "data" / f"champions_{role}.txt"
    if from_csv and csv_path.exists():
        rows = list(csv.DictReader(csv_path.open(encoding="utf-8")))
        for r in rows:
            r["games"] = int(r["games"]) if r["games"] else 0
        champs = [c for c in champs_txt.read_text(encoding="utf-8").splitlines() if c]
        print(f"== {role}: rebuilding from {csv_path.name} ({len(rows)} rows, {len(champs)} champions)")
    else:
        print(f"== {role}: discovering champion pool")
        champs = sorted(set(pool_for(role, delay)) | set(old["champs"]))
        print(f"   {len(champs)} champions")
        champs_txt.write_text("\n".join(champs) + "\n", encoding="utf-8")
        print(f"== {role}: fetching counters ({len(champs)} requests, {delay}s apart)")
        rows = fetch_all(role, champs, delay)
    champs = sorted(set(champs) | set(old["champs"]))
    comb = combined(rows)
    hand = role_data.CHAMPS.get(role, {})
    archs = role_data.ARCHETYPES.get(role, {})
    if old["champs"]:
        print(f"   refreshing existing {xlsx.name}: keeping {len(old['champs'])} champion rows' text, "
              f"{len(old['labels'])} hand labels and {len(old['tips'])} lane tips")

    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Champions"
    ws.append(["Champion", "Archetype", "Damage", "Strong in comps", "Good into", "Struggles into", "Pick when", "Blind-safe"])
    for c in champs:
        arch, dmg, comps, when, blind = hand.get(c, ("Flex", role_data.damage_of(c), "", "", "Mostly"))
        o = old["champs"].get(c)
        if o:
            if o.get("Archetype") == "Flex":  # placeholder in the workbook: the hand table wins
                o = {k: v for k, v in o.items() if k not in ("Archetype", "Damage", "Strong in comps", "Pick when", "Blind-safe") or v not in ("Flex", "Mostly", "", None)}
            arch, dmg, comps, when, blind = (o.get("Archetype") or arch, o.get("Damage") or dmg, o.get("Strong in comps") or comps,
                                             o.get("Pick when") or when, o.get("Blind-safe") or blind)
        words, g_default, b_default = archs.get(arch, ([], "", ""))
        pairs = sorted(((v["dnorm"], o_) for (cc, o_), v in comb.items() if cc == c and v["games"] >= MIN_GAMES), reverse=True)
        good = [o_ for d, o_ in pairs[:TOP_N] if d >= 1.0]
        bad = [o_ for d, o_ in pairs[::-1][:TOP_N] if d <= -1.0]
        good_txt = ", ".join(good) + (f"; {g_default}" if g_default else "")
        bad_txt = ", ".join(bad) + (f"; {b_default}" if b_default else "")
        if o:  # hand-written Good/Struggles into (no data list in front) wins; keep it after the data list
            for fld, data_list, default in (("Good into", good, g_default), ("Struggles into", bad, b_default)):
                txt = str(o.get(fld) or "").strip()
                if txt and not txt.startswith(", ".join(data_list)) and txt != default:
                    handtxt = txt.split("; ", 1)[-1] if data_list and txt.startswith(", ".join(data_list[:1])) else txt
                    merged = (", ".join(data_list) + "; " if data_list else "") + handtxt
                    if fld == "Good into": good_txt = merged
                    else: bad_txt = merged
        ws.append([c, arch, dmg, comps, good_txt, bad_txt, when, blind])

    ms = wb.create_sheet("Matchups")
    ms.append(["Champion", "Opponent", "Result for champion", "Lane tip", "Lola WR", "Lola dNorm", "Lola games",
               "Lola label", "Lola patch", "Lola source", "Mismatch"])
    n = 0
    keys = set(k for k in comb if k[0] in champs) | set(old["labels"]) | set(old["tips"])
    for (c, o) in sorted(keys):
        v = comb.get((c, o))
        handlab = old["labels"].get((c, o))
        tip = old["tips"].get((c, o), "")
        if v is None or (v["games"] < MIN_GAMES and not handlab and not tip):
            if not handlab and not tip:
                continue
        if v is None:
            ms.append([c, o, handlab or "", tip, None, None, None, "no data", "", "", ""])
        else:
            lab = lola.label(v["dnorm"], v["games"])
            mism = "yes" if handlab and norm_label(handlab) != norm_label(lab) and "low sample" not in lab else ""
            ms.append([c, o, handlab or lab, tip, round(v["wr"], 2), round(v["dnorm"], 2), v["games"], lab,
                       f"{v['patch']} {v['tier']}", v["source"], mism])
        n += 1
    cs = wb.create_sheet("Comps")
    if "Comps" in old["sheets"]:
        for r in old["sheets"]["Comps"]: cs.append(r)
    else:
        cs.append(["Style", "Words used in 'Strong in comps'"])
        for k, v in {"wombo": "wombo, teamfight", "poke": "poke, siege", "pick": "pick", "dive": "dive",
                     "split": "split, flank", "scaling": "scaling, late, front-to-back"}.items():
            cs.append([k, v])
    ns = wb.create_sheet("Notes")
    for r in old["sheets"].get("Notes", []):
        if not str(r[0] or "").startswith("Built "):
            ns.append(r)
    ns.append([f"Built {dt.date.today()} by build_role.py from Lolalytics Emerald+ (lane={role_data.ROLES[role]['lane']}) "
               f"and role_data.py. Champion text, hand labels and lane tips are kept across rebuilds; rows with a hand label "
               f"or tip are kept even without Lolalytics data."])
    from safe_save import save_workbook; save_workbook(wb, xlsx)
    print(f"== {role}: {xlsx} written: {len(champs)} champions, {n} matchup rows")


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    from_csv = "--from-csv" in sys.argv
    roles = args or ["all"]
    if roles == ["all"]:
        roles = list(role_data.ROLES)
    for r in roles:
        r = role_data.ROLE_ALIASES.get(r, r)
        build(r, delay=1.5, from_csv=from_csv)


if __name__ == "__main__":
    main()
