#!/usr/bin/env python3
"""Turn the Reddit thread dumps in data/reddit/ into per-matchup lane tips, deterministically (no LLM).

For every champion file, each paragraph of a thread body (and of fetched comments) that mentions another champion
becomes a candidate tip for (champion, opponent). Candidates are ranked by the thread's recency_weight and how
matchup-flavoured the paragraph is; paragraphs about items/runes from threads older than 2 years are dropped.

Usage:
  python3 extract_tips.py build [--reddit data/reddit] [--out data/reddit_tips.csv] [--per-pair 3]
  python3 extract_tips.py show Zed Viktor            # print every snippet for one pair, newest first
  python3 extract_tips.py apply [--xlsx midlane_overview.xlsx] [--csv data/reddit_tips.csv]
      adds "Reddit tips", "Reddit mentions" and "Reddit newest" columns to the Matchups sheet. Tips from the
      opponent's subreddit about this champion are included, prefixed with "[<opponent> mains]".
"""
import argparse
import csv
import json
import re
from collections import defaultdict
from pathlib import Path

CHAMPIONS = """Aatrox, Ahri, Akali, Akshan, Alistar, Ambessa, Amumu, Anivia, Annie, Aphelios, Ashe, Aurelion Sol,
Aurora, Azir, Bard, Bel'Veth, Blitzcrank, Brand, Braum, Briar, Caitlyn, Camille, Cassiopeia, Cho'Gath, Corki, Darius,
Diana, Dr. Mundo, Draven, Ekko, Elise, Evelynn, Ezreal, Fiddlesticks, Fiora, Fizz, Galio, Gangplank, Garen, Gnar,
Gragas, Graves, Gwen, Hecarim, Heimerdinger, Hwei, Illaoi, Irelia, Ivern, Janna, Jarvan IV, Jax, Jayce, Jhin, Jinx,
K'Sante, Kai'Sa, Kalista, Karma, Karthus, Kassadin, Katarina, Kayle, Kayn, Kennen, Kha'Zix, Kindred, Kled, Kog'Maw,
LeBlanc, Lee Sin, Leona, Lillia, Lissandra, Lucian, Lulu, Lux, Malphite, Malzahar, Maokai, Master Yi, Mel, Milio,
Miss Fortune, Mordekaiser, Morgana, Naafiri, Nami, Nasus, Nautilus, Neeko, Nidalee, Nilah, Nocturne, Nunu & Willump,
Olaf, Orianna, Ornn, Pantheon, Poppy, Pyke, Qiyana, Quinn, Rakan, Rammus, Rek'Sai, Rell, Renata Glasc, Renekton,
Rengar, Riven, Rumble, Ryze, Samira, Sejuani, Senna, Seraphine, Sett, Shaco, Shen, Shyvana, Singed, Sion, Sivir,
Skarner, Smolder, Sona, Soraka, Swain, Sylas, Syndra, Tahm Kench, Taliyah, Talon, Taric, Teemo, Thresh, Tristana,
Trundle, Tryndamere, Twisted Fate, Twitch, Udyr, Urgot, Varus, Vayne, Veigar, Vel'Koz, Vex, Vi, Viego, Viktor,
Vladimir, Volibear, Warwick, Wukong, Xayah, Xerath, Xin Zhao, Yasuo, Yone, Yorick, Yunara, Yuumi, Zac, Zed, Zeri,
Ziggs, Zilean, Zoe, Zyra"""
ALIASES = {
    "tf": "Twisted Fate", "asol": "Aurelion Sol", "aurelion": "Aurelion Sol", "lb": "LeBlanc", "leblanc": "LeBlanc",
    "kass": "Kassadin", "kat": "Katarina", "cass": "Cassiopeia", "cassio": "Cassiopeia", "vlad": "Vladimir",
    "ori": "Orianna", "malz": "Malzahar", "liss": "Lissandra", "yas": "Yasuo", "velkoz": "Vel'Koz", "kha": "Kha'Zix",
    "khazix": "Kha'Zix", "mundo": "Dr. Mundo", "j4": "Jarvan IV", "jarvan": "Jarvan IV", "yi": "Master Yi",
    "mf": "Miss Fortune", "nunu": "Nunu & Willump", "gp": "Gangplank", "trynd": "Tryndamere", "morde": "Mordekaiser",
    "kench": "Tahm Kench", "tahm": "Tahm Kench", "heimer": "Heimerdinger", "fiddle": "Fiddlesticks", "blitz": "Blitzcrank",
    "naut": "Nautilus", "renata": "Renata Glasc", "ksante": "K'Sante", "kaisa": "Kai'Sa", "chogath": "Cho'Gath",
    "cho": "Cho'Gath", "belveth": "Bel'Veth", "kog": "Kog'Maw", "kogmaw": "Kog'Maw", "reksai": "Rek'Sai",
    "xin": "Xin Zhao", "lee": "Lee Sin", "wukong": "Wukong", "panth": "Pantheon", "voli": "Volibear", "ww": "Warwick",
    "eve": "Evelynn", "nid": "Nidalee", "sera": "Seraphine", "trist": "Tristana", "cait": "Caitlyn", "aph": "Aphelios",
}
CASE_SENSITIVE = {"Vi", "Mel", "Zac", "Sett", "Bard", "Shen", "Sion", "Nami", "Lux", "Jax", "Ori", "Kat", "Cho", "Lee", "Yi", "Eve"}
MATCHUP_RE = re.compile(r"\b(match-?ups?|counter|lane|laning|poke|trade|trading|all-?in|roam|wave|freeze|prio|"
                        r"dodge|ban|skill|favou?red|easy|hard|free|unplayable|cs|farm|level|lvl|ult|flash|zone)\b", re.I)
ITEM_RE = re.compile(r"\b(items?|build|mythic|boots|runes?|keystone)\b", re.I)


def build_patterns():
    names = {n.strip(): n.strip() for n in CHAMPIONS.replace("\n", " ").split(",") if n.strip()}
    pats = []
    for n in names:
        pats.append((re.compile(r"(?<![A-Za-z'])" + re.escape(n) + r"(?![A-Za-z'])",
                                0 if n in CASE_SENSITIVE else re.I), n))
    for a, n in ALIASES.items():
        flags = 0 if a.capitalize() in CASE_SENSITIVE else re.I
        pats.append((re.compile(r"(?<![A-Za-z'])" + re.escape(a) + r"(?![A-Za-z'])", flags), n))
    return pats


PATTERNS = build_patterns()


def norm(n: str) -> str:
    return re.sub(r"[^a-z0-9]", "", n.lower())


def mentions(text: str):
    found = {}
    for pat, name in PATTERNS:
        if name not in found and pat.search(text):
            found[name] = True
    return list(found)


def paragraphs(t: dict):
    yield "post", t.get("title", "") + "\n" + t.get("body", "")
    for c in t.get("comments", []):
        yield "comment", c.get("body", "")


def snippet(par: str, pat_names, limit=600) -> str:
    par = par.strip()
    if len(par) <= limit:
        return par
    for pat, name in PATTERNS:
        if name in pat_names:
            m = pat.search(par)
            if m:
                a = max(0, m.start() - limit // 2)
                return ("..." if a else "") + par[a:a + limit].strip() + "..."
    return par[:limit] + "..."


def build(reddit_dir: Path, out: Path, per_pair: int):
    pairs = defaultdict(list)
    for f in sorted(reddit_dir.glob("*.json")):
        d = json.loads(f.read_text(encoding="utf-8"))
        champ = d["champion"]
        for t in d["threads"]:
            w = t.get("recency_weight", 0.5)
            stale = t.get("stale_items", False)
            title_opps = [o for o in mentions(t.get("title", "")) if norm(o) != norm(champ)]
            for kind, body in paragraphs(t):
                for par in re.split(r"\n+", body):
                    if len(par) < 40:
                        continue
                    if stale and ITEM_RE.search(par):
                        continue
                    opps = [o for o in mentions(par) if norm(o) != norm(champ)]
                    units = [(par, opps)]
                    if len(opps) > 4:
                        # crowded paragraph (tier list, AMA): attribute at sentence level with one sentence of context
                        sents = re.split(r"(?<=[.!?])\s+", par)
                        units = []
                        for i, sent in enumerate(sents):
                            so = [o for o in mentions(sent) if norm(o) != norm(champ)]
                            if so and len(so) <= 4:
                                units.append((" ".join(sents[i:i + 2])[:400], so))
                    elif not opps and title_opps and len(title_opps) <= 2:
                        # question-style thread ("Help vs Viktor"): body and comments are about the title's champion
                        units = [(par, title_opps)]
                    for text, unit_opps in units:
                        if not unit_opps:
                            continue
                        mu = len(MATCHUP_RE.findall(text))
                        for o in unit_opps:
                            pairs[(champ, o)].append({
                                "score": w * (1 + min(mu, 6) / 3) * (0.8 if kind == "comment" else 1),
                                "weight": w, "published": t["published"], "url": t["url"], "kind": kind,
                                "text": snippet(text, [o]),
                            })
    rows = []
    for (champ, opp), cands in sorted(pairs.items()):
        cands.sort(key=lambda c: (c["score"], c["published"]), reverse=True)
        top = cands[:per_pair]
        rows.append({
            "champion": champ, "opponent": opp, "mentions": len(cands),
            "weighted": round(sum(c["weight"] for c in cands), 1),
            "newest": max(c["published"] for c in cands),
            **{f"tip{i + 1}": c["text"] for i, c in enumerate(top)},
            **{f"src{i + 1}": f"{c['published']} {c['url']}" for i, c in enumerate(top)},
        })
    out.parent.mkdir(parents=True, exist_ok=True)
    fields = ["champion", "opponent", "mentions", "weighted", "newest"] + \
             [f"tip{i + 1}" for i in range(per_pair)] + [f"src{i + 1}" for i in range(per_pair)]
    with out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)
    print(f"{len(rows)} champion/opponent pairs from {len(list(reddit_dir.glob('*.json')))} files -> {out}")


def show(reddit_dir: Path, champ: str, opp: str):
    f = next((p for p in reddit_dir.glob("*.json") if norm(json.loads(p.read_text(encoding="utf-8"))["champion"]) == norm(champ)), None)
    if not f:
        raise SystemExit(f"no file for {champ} in {reddit_dir}")
    d = json.loads(f.read_text(encoding="utf-8"))
    hits = []
    for t in sorted(d["threads"], key=lambda t: t["published"], reverse=True):
        for kind, body in paragraphs(t):
            for par in re.split(r"\n+", body):
                if len(par) >= 40 and any(norm(o) == norm(opp) for o in mentions(par)):
                    hits.append((t["published"], t["recency_weight"], kind, t["title"], par.strip()))
    for pub, w, kind, title, par in hits:
        print(f"--- {pub} (w={w}, {kind}) {title[:70]}\n{par[:900]}\n")
    print(f"{len(hits)} snippets for {champ} vs {opp}")


def apply(xlsx: Path, csv_path: Path):
    import openpyxl
    tips = {}
    with csv_path.open(newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            tips[(norm(r["champion"]), norm(r["opponent"]))] = r
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["Matchups"]
    header = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]
    ci, oi = header.index("champion") + 1, header.index("opponent") + 1
    idx = {}
    for name in ("Reddit tips", "Reddit mentions", "Reddit newest"):
        if name.lower() in header:
            idx[name] = header.index(name.lower()) + 1
        else:
            ws.cell(row=1, column=ws.max_column + 1, value=name)
            header.append(name.lower()); idx[name] = len(header)
    n = 0
    for row in range(2, ws.max_row + 1):
        champ, opp = ws.cell(row, ci).value, ws.cell(row, oi).value
        if not champ or not opp:
            continue
        fwd, rev = tips.get((norm(champ), norm(opp))), tips.get((norm(opp), norm(champ)))
        parts, ment, newest = [], 0, ""
        for r, tag in ((fwd, ""), (rev, f"[{opp} mains] ")):
            if not r:
                continue
            ment += int(r["mentions"]); newest = max(newest, r["newest"])
            parts += [tag + r[k] for k in ("tip1", "tip2", "tip3") if r.get(k)]
        if parts:
            n += 1
            ws.cell(row, idx["Reddit tips"], value=" | ".join(parts)[:4000])
            ws.cell(row, idx["Reddit mentions"], value=ment)
            ws.cell(row, idx["Reddit newest"], value=newest)
    from safe_save import save_workbook; save_workbook(wb, xlsx)
    print(f"tips written for {n} of {ws.max_row - 1} matchup rows")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = p.add_subparsers(dest="cmd", required=True)
    b = sp.add_parser("build"); b.add_argument("--reddit", default="data/reddit"); b.add_argument("--out", default="data/reddit_tips.csv"); b.add_argument("--per-pair", type=int, default=3)
    s = sp.add_parser("show"); s.add_argument("champ"); s.add_argument("opp"); s.add_argument("--reddit", default="data/reddit")
    a = sp.add_parser("apply"); a.add_argument("--xlsx", default="midlane_overview.xlsx"); a.add_argument("--csv", default="data/reddit_tips.csv")
    args = p.parse_args()
    if args.cmd == "build":
        build(Path(args.reddit), Path(args.out), args.per_pair)
    elif args.cmd == "show":
        show(Path(args.reddit), args.champ, args.opp)
    else:
        apply(Path(args.xlsx), Path(args.csv))


if __name__ == "__main__":
    main()
